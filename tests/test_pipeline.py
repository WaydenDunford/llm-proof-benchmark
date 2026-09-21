"""Canonical proofs/evaluations and idempotent workbook integration tests."""
import json
from pathlib import Path
import shutil
from openpyxl import load_workbook
import pytest
from src.benchmark import run_benchmark
from src.dual import bundle, import_chatgpt, math_shepherd

REPO = Path(__file__).resolve().parents[1]


@pytest.fixture
def root(tmp_path):
    for name in ("configs", "prompts", "theorems"):
        shutil.copytree(REPO / name, tmp_path / name)
    return tmp_path


def chatgpt_file(root, proofs):
    data = json.loads((REPO / "tests/fixtures/sample_chatgpt_evaluation.json").read_text())
    data["run_id"] = proofs["run_id"]; data["theorem_id"] = proofs["theorem"]["id"]
    target = root / "chatgpt.json"; target.write_text(json.dumps(data), encoding="utf-8-sig")
    return target


def test_one_proofs_file_one_evaluations_file_and_excel(root):
    proof_file = run_benchmark(root, "example_theorem", backend="mock")
    proofs = json.loads(proof_file.read_text())
    assert proof_file.parent.name == "proofs"
    assert len(proofs["proofs"]) == 3
    assert len({proof["proof_id"] for proof in proofs["proofs"]}) == 3
    evaluation_file = math_shepherd(root, proof_file, "mock")
    evaluations = json.loads(evaluation_file.read_text())
    assert evaluation_file.parent.name == "evaluations"
    assert len(evaluations["evaluations"]) == 3
    assert all(item["math_shepherd"]["status"] == "success" for item in evaluations["evaluations"])
    assert all(item["chatgpt"]["status"] == "pending" for item in evaluations["evaluations"])
    book = load_workbook(root / "results/benchmark_results.xlsx")
    assert book.sheetnames == ["Results", "Proofs", "MathShepherdSteps", "Runs", "Summary"]
    assert book["Results"].max_row == 4 and book["Proofs"].max_row == 4
    assert book["MathShepherdSteps"].max_row > 1


def test_blind_bundle_and_import_updates_same_file_and_excel(root):
    proof_file = run_benchmark(root, "example_theorem", backend="mock")
    original = math_shepherd(root, proof_file, "mock")
    markdown, mapping = bundle(root, proof_file)
    text = markdown.read_text(); proofs = json.loads(proof_file.read_text())
    assert mapping.exists()
    for proof in proofs["proofs"]:
        assert proof["model_name"] not in text and proof["model_id"] not in text
        assert str(proof["runtime_seconds"]) not in text
    target = import_chatgpt(root, proofs["run_id"], chatgpt_file(root, proofs))
    assert target == original
    evaluations = json.loads(target.read_text())
    assert all(item["chatgpt"]["status"] == "success" for item in evaluations["evaluations"])
    book = load_workbook(root / "results/benchmark_results.xlsx")
    rows_before = book["Results"].max_row
    import_chatgpt(root, proofs["run_id"], chatgpt_file(root, proofs))
    book = load_workbook(root / "results/benchmark_results.xlsx")
    assert book["Results"].max_row == rows_before
    assert [row[23].value for row in book["Results"].iter_rows(min_row=2) if row[0].value == proofs["run_id"]] == [9, 9, 9]


def test_historical_runs_are_preserved(root):
    first = run_benchmark(root, "example_theorem", backend="mock"); math_shepherd(root, first, "mock")
    second = run_benchmark(root, "T002", backend="mock"); math_shepherd(root, second, "mock")
    book = load_workbook(root / "results/benchmark_results.xlsx")
    run_ids = {row[0].value for row in book["Runs"].iter_rows(min_row=2)}
    assert len(run_ids) == 2


def test_import_rejects_missing_proof(root):
    proof_file = run_benchmark(root, "example_theorem", backend="mock"); math_shepherd(root, proof_file, "mock"); bundle(root, proof_file)
    proofs = json.loads(proof_file.read_text()); data = json.loads(chatgpt_file(root, proofs).read_text(encoding="utf-8-sig")); data["evaluations"].pop()
    bad = root / "bad.json"; bad.write_text(json.dumps(data))
    with pytest.raises(ValueError, match="proof IDs"):
        import_chatgpt(root, proofs["run_id"], bad)
